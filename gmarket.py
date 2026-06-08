import os
import re
import importlib
import ssl
import subprocess
import time
import urllib.error
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font
import undetected_chromedriver as uc
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from urllib.parse import quote


def normalize(text):
    return re.sub(r"\s+", "", text or "")


def normalize_match(text):
    return normalize(text).casefold()


DEFAULT_OUTPUT_WORKBOOK = Path(__file__).with_name("빅스_판매량_20260604.xlsx")
OUTPUT_WORKBOOK = DEFAULT_OUTPUT_WORKBOOK
GMARKET_HOME_URL = "https://www.gmarket.co.kr"
HEADER_ROW = 3
SELLER_COLUMN = 3
PRODUCT_NAME_COLUMN = 4
FIRST_DATE_COLUMN = 6
FIRST_PRODUCT_ROW = 4
MAX_PRODUCTS = 999
FIRST_PRICE_COLUMN = 39
PRICE_GROUP_WIDTH = 3


def configure_ssl_certificates():
    try:
        truststore = importlib.import_module("truststore")
        truststore.inject_into_ssl()
        return "truststore"
    except Exception as truststore_error:
        try:
            certifi = importlib.import_module("certifi")
            os.environ.setdefault("SSL_CERT_FILE", certifi.where())
            os.environ.setdefault("REQUESTS_CA_BUNDLE", certifi.where())
            return "certifi"
        except Exception:
            print(f"SSL certificate setup skipped: {truststore_error}")
            return None


def is_ssl_certificate_error(error):
    seen = set()
    pending = [error]

    while pending:
        current = pending.pop()
        if current is None or id(current) in seen:
            continue

        seen.add(id(current))
        if isinstance(current, ssl.SSLCertVerificationError):
            return True

        if isinstance(current, urllib.error.URLError):
            pending.append(current.reason)

        message = str(current)
        if "CERTIFICATE_VERIFY_FAILED" in message or "certificate verify failed" in message:
            return True

        pending.append(getattr(current, "__cause__", None))
        pending.append(getattr(current, "__context__", None))

    return False


def parse_chrome_major_version(text):
    match = re.search(r"(\d+)\.", text or "")
    if not match:
        return None
    return int(match.group(1))


def get_windows_file_major_version(file_path):
    if os.name != "nt":
        return None

    try:
        import ctypes

        version_dll = ctypes.WinDLL("version", use_last_error=True)
        file_path_text = str(file_path)
        size = version_dll.GetFileVersionInfoSizeW(file_path_text, None)
        if not size:
            return None

        buffer = ctypes.create_string_buffer(size)
        if not version_dll.GetFileVersionInfoW(file_path_text, 0, size, buffer):
            return None

        value = ctypes.c_void_p()
        value_size = ctypes.c_uint()
        if not version_dll.VerQueryValueW(
            buffer,
            "\\",
            ctypes.byref(value),
            ctypes.byref(value_size),
        ):
            return None

        class VS_FIXEDFILEINFO(ctypes.Structure):
            _fields_ = [
                ("dwSignature", ctypes.c_uint32),
                ("dwStrucVersion", ctypes.c_uint32),
                ("dwFileVersionMS", ctypes.c_uint32),
                ("dwFileVersionLS", ctypes.c_uint32),
                ("dwProductVersionMS", ctypes.c_uint32),
                ("dwProductVersionLS", ctypes.c_uint32),
                ("dwFileFlagsMask", ctypes.c_uint32),
                ("dwFileFlags", ctypes.c_uint32),
                ("dwFileOS", ctypes.c_uint32),
                ("dwFileType", ctypes.c_uint32),
                ("dwFileSubtype", ctypes.c_uint32),
                ("dwFileDateMS", ctypes.c_uint32),
                ("dwFileDateLS", ctypes.c_uint32),
            ]

        fixed_info = ctypes.cast(
            value,
            ctypes.POINTER(VS_FIXEDFILEINFO),
        ).contents
        return fixed_info.dwFileVersionMS >> 16
    except Exception:
        return None


def get_chrome_executable_candidates():
    candidates = [
        Path(r"C:\Program Files\Google\Chrome\Application\chrome.exe"),
        Path(r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe"),
    ]

    local_app_data = os.environ.get("LOCALAPPDATA")
    if local_app_data:
        candidates.append(
            Path(local_app_data) / "Google" / "Chrome" / "Application" / "chrome.exe"
        )

    return candidates


def get_chrome_installation():
    for chrome_path in get_chrome_executable_candidates():
        if not chrome_path.exists():
            continue

        try:
            result = subprocess.run(
                [str(chrome_path), "--version"],
                check=False,
                capture_output=True,
                text=True,
                timeout=10,
            )
        except Exception:
            continue

        version = parse_chrome_major_version(result.stdout or result.stderr)
        if not version:
            version = get_windows_file_major_version(chrome_path)

        if version:
            return version, chrome_path

    return None


def get_chrome_major_version():
    chrome_installation = get_chrome_installation()
    if not chrome_installation:
        return None

    chrome_major_version, _chrome_path = chrome_installation
    return chrome_major_version


def create_undetected_chrome(**kwargs):
    try:
        return uc.Chrome(**kwargs)
    except Exception as error:
        if is_ssl_certificate_error(error):
            print(
                "ChromeDriver download failed because SSL certificate verification failed. "
                "Run Windows Update and try again, or rebuild the installer with updated "
                "certificate support."
            )
        raise


def create_chrome_driver(options):
    configure_ssl_certificates()
    chrome_installation = get_chrome_installation()
    if chrome_installation:
        chrome_major_version, chrome_path = chrome_installation
        print(
            f"설치된 Chrome {chrome_major_version} 버전으로 실행합니다: "
            f"{chrome_path}"
        )
        return create_undetected_chrome(
            version_main=chrome_major_version,
            browser_executable_path=str(chrome_path),
            options=options,
            use_subprocess=True,
        )

    return create_undetected_chrome(options=options, use_subprocess=True)


def get_tracking_sheet(wb):
    for ws in wb.worksheets:
        seller_header = ws.cell(2, SELLER_COLUMN).value
        product_header = ws.cell(2, PRODUCT_NAME_COLUMN).value
        if seller_header == "판매자" and product_header == "상품명":
            return ws

    return wb.active


def load_products_from_workbook(workbook_path=DEFAULT_OUTPUT_WORKBOOK):
    wb = load_workbook(workbook_path, data_only=True)
    ws = get_tracking_sheet(wb)
    products = []

    for row in range(FIRST_PRODUCT_ROW, ws.max_row + 1):
        seller_name = ws.cell(row, SELLER_COLUMN).value
        product_name = ws.cell(row, PRODUCT_NAME_COLUMN).value

        if not seller_name or not product_name:
            continue

        products.append({
            "상품명": str(product_name).strip(),
            "판매자": str(seller_name).strip(),
            "엑셀행": row,
        })

        if len(products) >= MAX_PRODUCTS:
            break

    return products


def get_date_label():
    now = datetime.now()
    return f"{now.month}/{now.day}\n{(now.hour % 12 or 12)}{'AM' if now.hour < 12 else 'PM'}"


def get_price_label():
    now = datetime.now()
    return f"{now.month}/{now.day}\n 가격"


def get_reward_label():
    now = datetime.now()
    return f"{now.month}/{now.day}\n적립금"


def same_date_label(left, right):
    return normalize_match(str(left or "")) == normalize_match(right)


def get_month_day_key(value):
    match = re.match(r"\s*(\d{1,2})/(\d{1,2})", str(value or ""))
    if not match:
        return None
    return int(match.group(1)), int(match.group(2))


def same_sales_date_label(left, right):
    left_key = get_month_day_key(left)
    right_key = get_month_day_key(right)
    if left_key and right_key:
        return left_key == right_key
    return same_date_label(left, right)


def find_or_create_date_column(ws, date_label):
    for col in range(FIRST_DATE_COLUMN, ws.max_column + 1):
        if same_sales_date_label(ws.cell(HEADER_ROW, col).value, date_label):
            ws.cell(HEADER_ROW, col).value = date_label
            return col

    for col in range(FIRST_DATE_COLUMN, ws.max_column + 2):
        value = ws.cell(HEADER_ROW, col).value
        if value is None or str(value).strip() == "":
            ws.cell(HEADER_ROW, col).value = date_label
            return col

    raise RuntimeError("날짜를 입력할 빈 컬럼을 찾지 못했습니다.")


def coerce_count(value):
    if isinstance(value, str) and value.isdigit():
        return int(value)
    return value


def coerce_amount(value):
    if isinstance(value, str) and value.isdigit():
        return int(value)
    return value


def price_reward_group_is_empty(ws, price_col, reward_col, results):
    for result in results:
        row = result.get("엑셀행")
        if not row:
            continue
        price_value = ws.cell(row, price_col).value
        reward_value = ws.cell(row, reward_col).value
        if price_value not in (None, "") or reward_value not in (None, ""):
            return False
    return True


def find_or_create_price_reward_columns(ws, results):
    price_label = get_price_label()
    reward_label = get_reward_label()

    for price_col in range(FIRST_PRICE_COLUMN, ws.max_column + PRICE_GROUP_WIDTH + 1, PRICE_GROUP_WIDTH):
        reward_col = price_col + 1
        price_header = ws.cell(HEADER_ROW, price_col).value
        reward_header = ws.cell(HEADER_ROW, reward_col).value

        if (
            same_date_label(price_header, price_label)
            and same_date_label(reward_header, reward_label)
        ):
            return price_col, reward_col

    for price_col in range(FIRST_PRICE_COLUMN, ws.max_column + PRICE_GROUP_WIDTH + 1, PRICE_GROUP_WIDTH):
        reward_col = price_col + 1
        price_header = ws.cell(HEADER_ROW, price_col).value
        reward_header = ws.cell(HEADER_ROW, reward_col).value

        if (
            price_header in (None, "")
            and reward_header in (None, "")
            and price_reward_group_is_empty(ws, price_col, reward_col, results)
        ):
            ws.cell(HEADER_ROW, price_col).value = price_label
            ws.cell(HEADER_ROW, reward_col).value = reward_label
            return price_col, reward_col

    raise RuntimeError("가격/적립금을 입력할 빈 컬럼을 찾지 못했습니다.")


def save_results_to_workbook(results, workbook_path=DEFAULT_OUTPUT_WORKBOOK):
    wb = load_workbook(workbook_path)
    ws = get_tracking_sheet(wb)

    date_label = get_date_label()
    date_col = find_or_create_date_column(ws, date_label)
    price_col, reward_col = find_or_create_price_reward_columns(ws, results)

    saved = 0
    price_saved = 0
    skipped = []

    for result in results:
        row = result.get("엑셀행")

        if not row:
            skipped.append(result["상품명"])
            continue

        ws.cell(row, date_col).value = coerce_count(result["누적판매량"])
        ws.cell(row, price_col).value = coerce_amount(result["가격"])
        ws.cell(row, reward_col).value = coerce_amount(result["적립금"])
        saved += 1
        price_saved += 1

    wb.save(workbook_path)

    print(f"\n저장 완료: {workbook_path}")
    print(f"저장 위치: {ws.cell(HEADER_ROW, date_col).coordinate} 날짜 컬럼")
    print(
        "가격/적립금 저장 위치: "
        f"{ws.cell(HEADER_ROW, price_col).coordinate}/"
        f"{ws.cell(HEADER_ROW, reward_col).coordinate}"
    )
    print(f"저장 건수: {saved}건")
    print(f"가격/적립금 저장 건수: {price_saved}건")

    if skipped:
        print(f"엑셀 행 정보 없음으로 건너뜀: {len(skipped)}건")
        for name in skipped:
            print(f" - {name}")


def get_seller_name(item):
    seller_texts = item.find_elements(By.CSS_SELECTOR, "span.text__seller")
    if seller_texts:
        text = seller_texts[0].text.strip()
        if text:
            return text

    imgs = item.find_elements(
        By.CSS_SELECTOR,
        ".box__information_seller a.link__shop img.image__logo"
    )

    for img in imgs:
        alt = img.get_attribute("alt")
        if alt:
            return alt.strip()

    return ""


def digits_only(text):
    value = "".join(filter(str.isdigit, text or ""))
    return value if value else "표시없음"


def get_first_text_by_selectors(item, selectors):
    for selector in selectors:
        elements = item.find_elements(By.CSS_SELECTOR, selector)
        for element in elements:
            text = element.text.strip()
            if not text:
                text = (element.get_attribute("innerText") or "").strip()
            if text:
                return text
    return ""


def get_price(item):
    text = get_first_text_by_selectors(
        item,
        [
            ".box__price-seller strong.text__value",
            ".box__price strong.text__value",
            ".box__price-seller .text__value",
            ".box__price .text__value",
            "strong.text__value",
        ],
    )
    return digits_only(text)


def get_reward(item):
    item_text = item.text
    patterns = [
        r"([\d,]+)\s*원?\s*적립",
        r"적립\s*([\d,]+)",
    ]

    for pattern in patterns:
        match = re.search(pattern, item_text)
        if match:
            return digits_only(match.group(1))

    text = get_first_text_by_selectors(
        item,
        [
            ".box__benefit",
            ".box__item-benefit",
            ".list-item__benefit",
            ".box__coupon",
        ],
    )
    return digits_only(text) if "적립" in text else "표시없음"


def search_from_homepage(driver, keyword):
    driver.get(GMARKET_HOME_URL)
    search_input = WebDriverWait(driver, 15).until(
        EC.element_to_be_clickable(
            (By.CSS_SELECTOR, "#form__search-keyword, input[name='keyword']")
        )
    )
    search_input.clear()
    driver.execute_script(
        """
        arguments[0].value = arguments[1];
        arguments[0].dispatchEvent(new Event('input', {bubbles: true}));
        arguments[0].dispatchEvent(new Event('change', {bubbles: true}));
        arguments[0].closest('form').submit();
        """,
        search_input,
        keyword,
    )
    time.sleep(2)


def open_search_page(driver, keyword, use_homepage=False):
    if use_homepage:
        search_from_homepage(driver, keyword)
        return

    search_url = f"https://www.gmarket.co.kr/n/search?keyword={quote(keyword)}"
    driver.get(search_url)


def main(workbook_path=DEFAULT_OUTPUT_WORKBOOK):
    products = load_products_from_workbook(workbook_path)
    print(f"엑셀에서 상품 {len(products)}건을 읽었습니다.")

    options = uc.ChromeOptions()
    options.add_argument("--start-maximized")
    options.add_argument("--no-sandbox")
    options.add_argument("--disable-dev-shm-usage")
    options.add_argument("--disable-gpu")

    driver = create_chrome_driver(options)

    results = []

    try:
        for index, p in enumerate(products):
            raw_name = p["상품명"]
            seller_target = p["판매자"]

            keyword = re.sub(r"\([^)]*\)", "", raw_name).strip()

            print(f"\n검색 중: {raw_name}")
            open_search_page(driver, keyword, use_homepage=index == 0)

            try:
                WebDriverWait(driver, 10).until(
                    EC.presence_of_element_located((By.CSS_SELECTOR, "div.box__item-container"))
                )

            except:
                print("검색 결과 로딩 실패")
                continue

            items = driver.find_elements(By.CSS_SELECTOR, "div.box__item-container")
            found = False

            for item in items:
                try:
                    title = item.find_element(
                        By.CSS_SELECTOR,
                        "span.text__item"
                    ).text.strip()

                    seller = get_seller_name(item)

                    if (
                        normalize(keyword) in normalize(title)
                        and (
                            not seller_target
                            or normalize(seller_target) in normalize(seller)
                        )
                    ):
                        buy_count = "표시없음"

                        buy_els = item.find_elements(
                            By.CSS_SELECTOR,
                            ".list-item__pay-count, .text__buy-count"
                        )
                        if buy_els:
                            buy_count = "".join(
                                filter(str.isdigit, buy_els[0].text)
                            )

                        price = get_price(item)
                        reward = get_reward(item)

                        results.append({
                            "상품명": raw_name,
                            "판매자": seller_target,
                            "엑셀행": p.get("엑셀행"),
                            "누적판매량": buy_count,
                            "가격": price,
                            "적립금": reward,
                        })

                        print(
                            f"매칭 성공 → 판매자: {seller} / 판매량: {buy_count} "
                            f"/ 가격: {price} / 적립금: {reward}"
                        )
                        found = True
                        break

                except Exception:
                    continue

            if not found:
                results.append({
                    "상품명": raw_name,
                    "판매자": seller_target if seller_target else "관계없음",
                    "엑셀행": p.get("엑셀행"),
                    "누적판매량": "미발견",
                    "가격": "미발견",
                    "적립금": "미발견",
                })
                print("매칭 실패")

    finally:
        pass
        # driver.quit()   # 브라우저 자동 종료 원하면 활성화

    save_results_to_workbook(results, workbook_path)


if __name__ == "__main__":
    main()
